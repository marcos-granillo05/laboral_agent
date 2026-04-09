import json
import os
import uuid
import traceback
from io import BytesIO
from datetime import datetime, timezone, timedelta

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from google.cloud import storage
from google.oauth2 import service_account
from google.genai import types
from google.adk.agents.callback_context import CallbackContext


def dbg(step, **kwargs):
    print(f"\n[DEBUG] {step}")
    if kwargs:
        for k, v in kwargs.items():
            try:
                print(f"    - {k}: {v}")
            except Exception as e:
                print(f"    - {k}: <error printing value: {e}>")


# Para que no truene en dev ni prod je
def build_storage_client():
    dbg("build_storage_client:start")
    key_path = os.environ.get("ADK_CREDENTIALS_PATH")
    dbg("build_storage_client:key_path", key_path=key_path, exists=os.path.exists(key_path) if key_path else False)

    if key_path and os.path.exists(key_path):
        dbg("build_storage_client:using_service_account_file")
        creds = service_account.Credentials.from_service_account_file(key_path)
        client = storage.Client(credentials=creds, project=creds.project_id)
        dbg("build_storage_client:client_created_with_file", project_id=creds.project_id)
        return client, creds

    dbg("build_storage_client:using_default_credentials")
    client = storage.Client()
    dbg("build_storage_client:client_created_with_adc")
    return client, None


# def upload_to_gcs_public(bucket_name: str, object_name: str, data: bytes) -> dict:
    # dbg("upload_to_gcs_public:start", bucket_name=bucket_name, object_name=object_name, data_len=len(data) if data else 0)

    # client, creds = build_storage_client()
    # dbg("upload_to_gcs_public:client_ready", creds_present=creds is not None)

    # bucket = client.bucket(bucket_name)
    # blob = bucket.blob(object_name)
    # dbg("upload_to_gcs_public:bucket_blob_created")

    # try:
    #     blob.upload_from_string(
    #         data,
    #         content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    #     )
    #     dbg("upload_to_gcs_public:upload_ok")
    # except Exception as e:
    #     dbg("upload_to_gcs_public:upload_error", error=str(e), traceback=traceback.format_exc()) 
    #     raise

    # try:
    #     if creds is not None:
    #         dbg("upload_to_gcs_public:generating_signed_url_with_explicit_creds")
    #         url = blob.generate_signed_url(
    #             version="v4",
    #             expiration=timedelta(minutes=15),
    #             method="GET",
    #             credentials=creds
    #         )
    #     else:
    #         dbg("upload_to_gcs_public:generating_signed_url_with_default_env")
    #         url = blob.generate_signed_url(
    #             version="v4",
    #             expiration=timedelta(minutes=15),
    #             method="GET"
    #         )

    #     dbg("upload_to_gcs_public:signed_url_ok", url=url)
    # except Exception as e:
    #     dbg("upload_to_gcs_public:signed_url_error", error=str(e), traceback=traceback.format_exc()) # ACÁ SI DIO ERROR JE
    #     raise

    # return {
    #     "bucket_name": bucket_name,
    #     "object_name": object_name,
    #     "signed_url": url
    # }

def upload_to_gcs(bucket_name: str, object_name: str, data: bytes) -> dict:
    client, _ = build_storage_client()
    bucket = client.bucket(bucket_name)
    blob = bucket.blob(object_name)

    blob.upload_from_string(
        data,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    print(f"DEBUG: upload ok bucket={bucket_name} object={object_name}")

    return {
        "bucket_name": bucket_name,
        "object_name": object_name,
    }


def _apply_font(run, font_name="Arial", size_pt=None, color_rgb=(0, 0, 0), bold=False):
    dbg("_apply_font:start", text=run.text if run else None, font_name=font_name, size_pt=size_pt, color_rgb=color_rgb, bold=bold)

    run.font.name = font_name

    # Asegurar que exista rPr
    rPr = run._element.get_or_add_rPr()

    # Asegurar que exista rFonts
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)

    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)

    run.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)
    if size_pt is not None:
        run.font.size = Pt(size_pt)

    dbg("_apply_font:end", text=run.text if run else None)


def _set_default_font(doc: Document, font_name="Arial", size_pt=12):
    dbg("_set_default_font:start", font_name=font_name, size_pt=size_pt)
    style = doc.styles['Normal']
    style.font.name = font_name
    style.font.size = Pt(size_pt)

    rPr = style.element.get_or_add_rPr()
    rfonts = rPr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rPr.append(rfonts)

    rfonts.set(qn('w:ascii'), font_name)
    rfonts.set(qn('w:hAnsi'), font_name)
    rfonts.set(qn('w:cs'), font_name)
    dbg("_set_default_font:end")


def _add_horizontal_rule(paragraph, size="18", color="000000", space="1"):
    dbg("_add_horizontal_rule:start", size=size, color=color, space=space)

    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)

    pBdr.append(bottom)
    pPr.append(pBdr)

    dbg("_add_horizontal_rule:end")


def _fmt_fecha_from_ts(ts):
    dbg("_fmt_fecha_from_ts:start", ts=ts, ts_type=type(ts).__name__ if ts is not None else None)

    if ts is None:
        dbg("_fmt_fecha_from_ts:none")
        return None

    if isinstance(ts, str):
        try:
            dt = datetime.fromisoformat(ts.replace("Z", "+00:00"))
            dbg("_fmt_fecha_from_ts:parsed_iso", dt=dt.isoformat())
        except ValueError:
            dbg("_fmt_fecha_from_ts:iso_failed_try_epoch_string", ts=ts)
            ts_num = float(ts)
            dt = _dt_from_epoch(ts_num)
    else:
        dt = _dt_from_epoch(float(ts))

    meses = [
        "enero", "febrero", "marzo", "abril", "mayo", "junio",
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"
    ]
    result = f"{dt.day} de {meses[dt.month-1]} de {dt.year}"
    dbg("_fmt_fecha_from_ts:end", result=result)
    return result


def _dt_from_epoch(ts_num: float) -> datetime:
    dbg("_dt_from_epoch:start", ts_num=ts_num)
    if ts_num > 1e12:  # ms
        ts_num = ts_num / 1000.0
        dbg("_dt_from_epoch:converted_ms_to_s", converted=ts_num)

    dt = datetime.fromtimestamp(ts_num, tz=timezone.utc)
    dbg("_dt_from_epoch:end", dt=dt.isoformat())
    return dt


def _normalize_content(value):
    dbg("_normalize_content:start", value_type=type(value).__name__ if value is not None else None, value=value)

    if value is None:
        dbg("_normalize_content:none")
        return ""
    if isinstance(value, list):
        result = [str(x) for x in value if str(x).strip()]
        dbg("_normalize_content:list", result=result)
        return result
    if isinstance(value, str):
        result = value.strip()
        dbg("_normalize_content:str", result=result)
        return result

    result = str(value).strip()
    dbg("_normalize_content:other", result=result)
    return result


def _add_section(doc: Document, title: str, content, font_name="Arial"):
    dbg("_add_section:start", title=title, content_type=type(content).__name__ if content is not None else None)

    # Título
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = h.add_run(title)
    _apply_font(r, font_name=font_name, size_pt=12, bold=True)

    content_norm = _normalize_content(content)
    dbg("_add_section:content_normalized", title=title, content_norm=content_norm)

    # Contenido
    if isinstance(content_norm, list):
        for idx, item in enumerate(content_norm):
            dbg("_add_section:list_item", idx=idx, item=item)
            p = doc.add_paragraph(item, style="List Bullet")
            if p.runs:
                _apply_font(p.runs[0], font_name=font_name, size_pt=11)
    else:
        chunks = [c.strip() for c in content_norm.split("\n\n") if c.strip()]
        if not chunks:
            chunks = [""]
        dbg("_add_section:chunks", title=title, chunks=chunks)

        for idx, ch in enumerate(chunks):
            dbg("_add_section:chunk", idx=idx, chunk=ch)
            p = doc.add_paragraph(ch)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if p.runs:
                _apply_font(p.runs[0], font_name=font_name, size_pt=11)

    doc.add_paragraph()
    dbg("_add_section:end", title=title)


async def document_maker(analysis_json: str, tool_context: CallbackContext) -> dict:
    dbg("document_maker:start", analysis_json_type=type(analysis_json).__name__)

    try:
        dbg("document_maker:loading_json")
        analysis = json.loads(analysis_json) if isinstance(analysis_json, str) else analysis_json
        dbg("document_maker:json_loaded", analysis=analysis)

        dbg("document_maker:create_document")
        doc = Document()

        dbg("document_maker:set_margins")
        section = doc.sections[0]
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        font_name = "Arial"
        dbg("document_maker:set_default_font", font_name=font_name)
        _set_default_font(doc, font_name=font_name, size_pt=12)

        dbg("document_maker:add_header")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)

        for idx, line in enumerate([
            "PROCURADURÍA GENERAL DE LA REPÚBLICA",
            "UNIDAD PARA LA DEFENSA DE LOS DERECHOS",
            "LABORALES DE LA PERSONA TRABAJADORA",
            "PROCURADURÍA AUXILIAR DE xxxxxxxxxx",
            "Teléfono xxxx-xxxx",
        ]):
            dbg("document_maker:add_header_line", idx=idx, line=line)
            r = p.add_run(line + ("\n" if line != "Teléfono xxxx-xxxx" else ""))
            _apply_font(r, font_name=font_name, size_pt=11, bold=True)

        dbg("document_maker:add_rule")
        line_p = doc.add_paragraph()
        _add_horizontal_rule(line_p, size="18")
        line_p.paragraph_format.space_after = Pt(12)

        dbg("document_maker:format_date")
        fecha_str = _fmt_fecha_from_ts(analysis.get("timestamp"))
        dbg("document_maker:formatted_date", fecha_str=fecha_str)

        if fecha_str:
            dbg("document_maker:add_date_paragraph")
            pf = doc.add_paragraph(fecha_str)
            pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if pf.runs:
                _apply_font(pf.runs[0], font_name=font_name, size_pt=11)
            pf.paragraph_format.space_after = Pt(16)

         # claims_statement_facts: str = Field(...,description="Identificación de pretensiones sobre relato de hecho")   

        dbg("document_maker:extract_fields")
        personal_data = analysis.get("personal_data_summary", "")
        summary = analysis.get("summary", "")
        justification = analysis.get("justification", "")
        #suggestions = analysis.get("suggestions", "")
        claims_statement_facts  =analysis.get("claims_statement_facts", "") 
        user_dui = analysis.get("personal_data", {}).get("dui", "")

        dbg(
            "document_maker:fields_extracted",
            personal_data=personal_data,
            summary=summary,
            justification=justification,
       #     suggestions=suggestions,
            user_dui=user_dui,
            claims_statement_facts=claims_statement_facts
        )

        dbg("document_maker:add_sections")
        _add_section(doc, "Datos Personales", personal_data, font_name=font_name)
        _add_section(doc, "I. Resumen del relato del hecho", summary, font_name=font_name)
        _add_section(doc, "II. Trámite Asociado", claims_statement_facts,font_name=font_name)
        _add_section(doc, "III. Análisis Jurídico", justification, font_name=font_name)
      #  _add_section(doc, "IV. Ruta de Atención y Siguientes Pasos ", suggestions, font_name=font_name)

        dbg("document_maker:add_signature")
        atent = doc.add_paragraph("Atentamente,")
        atent.paragraph_format.space_after = Pt(28)
        if atent.runs:
            _apply_font(atent.runs[0], font_name=font_name, size_pt=11)

        firma = doc.add_paragraph()
        firma.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = firma.add_run("______________________________\n")
        _apply_font(r1, font_name=font_name, size_pt=11)
        r2 = firma.add_run("")
        _apply_font(r2, font_name=font_name, size_pt=11)

        dbg("document_maker:add_footer")
        footer_para = doc.sections[0].footer.paragraphs[0]
        footer_text = f'Generado el {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}'
        footer_para.text = footer_text
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if footer_para.runs:
            _apply_font(footer_para.runs[0], font_name=font_name, size_pt=8, color_rgb=(128, 128, 128))
        dbg("document_maker:footer_added", footer_text=footer_text)

        dbg("document_maker:save_doc_to_buffer")
        buffer = BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.seek(0)
        mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        dbg("document_maker:buffer_ready", bytes_len=len(docx_bytes), mime_type=mime_type)

        dbg("document_maker:get_bucket_env")
        bucket_name = os.environ["GCS_BUCKET_LABORAL"]
        dbg("document_maker:bucket_env_ok", bucket_name=bucket_name)

        case_id = uuid.uuid4().hex
        dbg("document_maker:case_id_generated", case_id=case_id)

        safe_dui = "".join(ch for ch in str(user_dui) if ch.isalnum()) or "sin_dui"
        dbg("document_maker:safe_dui", safe_dui=safe_dui)

        object_name = f"case_{safe_dui}/analytical_summary/{case_id}_{safe_dui}_analytical_summary_document.docx"
        output_filename = f"{case_id}_{safe_dui}_analytical_summary_document.docx"
        dbg("document_maker:file_names", object_name=object_name, output_filename=output_filename)

        dbg("document_maker:upload_to_gcs:start")
        save_gcp = upload_to_gcs(
            bucket_name=bucket_name,
            object_name=object_name,
            data=docx_bytes,
        )
        dbg("document_maker:upload_to_gcs:done", save_gcp=save_gcp)

        dbg("document_maker:create_artifact_part")
        artifact_part = types.Part(
            inline_data=types.Blob(data=docx_bytes, mime_type=mime_type)
        )
        dbg("document_maker:artifact_part_created")

        dbg("document_maker:save_artifact:start", output_filename=output_filename)
        version = await tool_context.save_artifact(
            filename=output_filename,
            artifact=artifact_part
        )
        dbg("document_maker:save_artifact:done", version=version)

        dbg("document_maker:success")
        return {
            "status": "ok",
            "message": f"El documento {output_filename} version {version} ha sido creado y disponible para descargar.",
            "gcs": save_gcp
        }

    except Exception as e:
        dbg("document_maker:EXCEPTION", error_type=type(e).__name__, error=str(e), traceback=traceback.format_exc())
        return {
            "status": "error",
            "error_type": type(e).__name__,
            "error_message": str(e),
            "traceback": traceback.format_exc()
        }


async def document_maker_V02(
    personal_data_summary: str,
    personal_data: dict,
    summary: str,
    justification: str,
    suggestions: str,
    legal_diagnosis: str = "",
    presuntion_status: str = "",
    actions: str = "",
    formats: str = "",
    timestamp=None,
    tool_context: CallbackContext = None,
) -> dict:
    dbg(
        "document_maker_V02:start",
        personal_data_summary_type=type(personal_data_summary).__name__,
        personal_data_type=type(personal_data).__name__,
        summary_type=type(summary).__name__,
        justification_type=type(justification).__name__,
        suggestions_type=type(suggestions).__name__,
    )

    try:
        analysis = {
            "personal_data_summary": personal_data_summary,
            "personal_data": personal_data,
            "summary": summary,
            "justification": justification,
            "suggestions": suggestions,
            "legal_diagnosis": legal_diagnosis,
            "presuntion_status": presuntion_status,
            "actions": actions,
            "formats": formats,
            "timestamp": timestamp,
        }
        dbg("document_maker_V02:analysis_built", analysis=analysis)

        dbg("document_maker_V02:create_document")
        doc = Document()

        dbg("document_maker_V02:set_margins")
        section = doc.sections[0]
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        font_name = "Arial"
        dbg("document_maker_V02:set_default_font", font_name=font_name)
        _set_default_font(doc, font_name=font_name, size_pt=12)

        dbg("document_maker_V02:add_header")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)

        for idx, line in enumerate([
            "PROCURADURÃA GENERAL DE LA REPÃšBLICA",
            "UNIDAD PARA LA DEFENSA DE LOS DERECHOS",
            "LABORALES DE LA PERSONA TRABAJADORA",
            "PROCURADURÃA AUXILIAR DE xxxxxxxxxx",
            "TelÃ©fono xxxx-xxxx",
        ]):
            dbg("document_maker_V02:add_header_line", idx=idx, line=line)
            r = p.add_run(line + ("\n" if line != "TelÃ©fono xxxx-xxxx" else ""))
            _apply_font(r, font_name=font_name, size_pt=11, bold=True)

        dbg("document_maker_V02:add_rule")
        line_p = doc.add_paragraph()
        _add_horizontal_rule(line_p, size="18")
        line_p.paragraph_format.space_after = Pt(12)

        dbg("document_maker_V02:format_date")
        fecha_str = _fmt_fecha_from_ts(analysis.get("timestamp"))
        dbg("document_maker_V02:formatted_date", fecha_str=fecha_str)

        if fecha_str:
            dbg("document_maker_V02:add_date_paragraph")
            pf = doc.add_paragraph(fecha_str)
            pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if pf.runs:
                _apply_font(pf.runs[0], font_name=font_name, size_pt=11)
            pf.paragraph_format.space_after = Pt(16)

        dbg("document_maker_V02:extract_fields")
        personal_data_block = analysis.get("personal_data_summary", "")
        summary_block = analysis.get("summary", "")
        justification_block = analysis.get("justification", "")
        suggestions_block = analysis.get("suggestions", "")
        user_dui = analysis.get("personal_data", {}).get("dui", "")

        dbg(
            "document_maker_V02:fields_extracted",
            personal_data=personal_data_block,
            summary=summary_block,
            justification=justification_block,
            suggestions=suggestions_block,
            user_dui=user_dui
        )

        dbg("document_maker_V02:add_sections")
        _add_section(doc, "Datos Personales", personal_data_block, font_name=font_name)
        _add_section(doc, "I. Resumen del relato del hecho", summary_block, font_name=font_name)
        _add_section(doc, "II. Leyes violadas", justification_block, font_name=font_name)
        _add_section(doc, "III. Consejos", suggestions_block, font_name=font_name)

        dbg("document_maker_V02:add_signature")
        atent = doc.add_paragraph("Atentamente,")
        atent.paragraph_format.space_after = Pt(28)
        if atent.runs:
            _apply_font(atent.runs[0], font_name=font_name, size_pt=11)

        firma = doc.add_paragraph()
        firma.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = firma.add_run("______________________________\n")
        _apply_font(r1, font_name=font_name, size_pt=11)
        r2 = firma.add_run("")
        _apply_font(r2, font_name=font_name, size_pt=11)

        dbg("document_maker_V02:add_footer")
        footer_para = doc.sections[0].footer.paragraphs[0]
        footer_text = f'Generado el {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}'
        footer_para.text = footer_text
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if footer_para.runs:
            _apply_font(footer_para.runs[0], font_name=font_name, size_pt=8, color_rgb=(128, 128, 128))
        dbg("document_maker_V02:footer_added", footer_text=footer_text)

        dbg("document_maker_V02:save_doc_to_buffer")
        buffer = BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.seek(0)
        mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        dbg("document_maker_V02:buffer_ready", bytes_len=len(docx_bytes), mime_type=mime_type)

        dbg("document_maker_V02:get_bucket_env")
        bucket_name = os.environ["GCS_BUCKET_LABORAL"]
        dbg("document_maker_V02:bucket_env_ok", bucket_name=bucket_name)

        case_id = uuid.uuid4().hex
        dbg("document_maker_V02:case_id_generated", case_id=case_id)

        safe_dui = "".join(ch for ch in str(user_dui) if ch.isalnum()) or "sin_dui"
        dbg("document_maker_V02:safe_dui", safe_dui=safe_dui)

        object_name = f"case_{safe_dui}/{case_id}_{safe_dui}_reporte_procesal.docx"
        output_filename = f"{case_id}_{safe_dui}_reporte_procesal.docx"
        dbg("document_maker_V02:file_names", object_name=object_name, output_filename=output_filename)

        dbg("document_maker_V02:upload_to_gcs:start")
        save_gcp = upload_to_gcs(
            bucket_name=bucket_name,
            object_name=object_name,
            data=docx_bytes,
        )
        dbg("document_maker_V02:upload_to_gcs:done", save_gcp=save_gcp)

        dbg("document_maker_V02:create_artifact_part")
        artifact_part = types.Part(
            inline_data=types.Blob(data=docx_bytes, mime_type=mime_type)
        )
        dbg("document_maker_V02:artifact_part_created")

        dbg("document_maker_V02:save_artifact:start", output_filename=output_filename)
        version = await tool_context.save_artifact(
            filename=output_filename,
            artifact=artifact_part
        )
        dbg("document_maker_V02:save_artifact:done", version=version)

        dbg("document_maker_V02:success")
        return {
            "status": "ok",
            "message": f"El documento {output_filename} version {version} ha sido creado y disponible para descargar.",
            "gcs": save_gcp
        }

    except Exception as e:
        dbg("document_maker_V02:EXCEPTION", error_type=type(e).__name__, error=str(e), traceback=traceback.format_exc())
        return {
            "status": "error",
            "error_type": type(e).__name__,
            "error_message": str(e),
            "traceback": traceback.format_exc()
        }
