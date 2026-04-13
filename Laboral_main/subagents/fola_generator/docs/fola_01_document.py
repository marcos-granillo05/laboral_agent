import json
import os
import uuid
import traceback
from io import BytesIO
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from google.cloud import storage
from google.oauth2 import service_account
from google.genai import types
from google.adk.agents.callback_context import CallbackContext


# =========================
# GCS
# =========================

def build_storage_client():
    key_path = os.environ.get("ADK_CREDENTIALS_PATH")

    if key_path and os.path.exists(key_path):
        creds = service_account.Credentials.from_service_account_file(key_path)
        client = storage.Client(credentials=creds, project=creds.project_id)
        return client, creds

    client = storage.Client()
    return client, None


def upload_to_gcs(bucket_name: str, object_name: str, data: bytes) -> dict:
    client, _ = build_storage_client()
    bucket = client.bucket(bucket_name)
    blob = bucket.blob(object_name)

    blob.upload_from_string(
        data,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    return {
        "bucket_name": bucket_name,
        "object_name": object_name,
    }


def _build_artifact_filename(case_id: str, safe_dui: str, prefix: str) -> str:
    short_case_id = case_id[:8]
    short_dui = safe_dui[:12]
    return f"{prefix}_{short_dui}_{short_case_id}.docx"


# =========================
# HELPERS DE FORMATO
# =========================

def _apply_font(run, font_name="Arial", size_pt=None, color_rgb=(0, 0, 0), bold=False):
    run.font.name = font_name

    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)

    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:cs"), font_name)

    run.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def _set_default_font(doc: Document, font_name="Arial", size_pt=11):
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(size_pt)

    rPr = style.element.get_or_add_rPr()
    rfonts = rPr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rPr.append(rfonts)

    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:cs"), font_name)


def _set_spacing(paragraph, before=0, after=0, line=1.0):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def _normalize_text(value):
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(x).strip() for x in value if str(x).strip())
    return str(value).strip()


def _safe(value, default=""):
    value = _normalize_text(value)
    return value if value else default


def _add_bottom_border(paragraph, size="6", color="000000", space="1"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)

    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_empty_line(doc, after=4):
    p = doc.add_paragraph()
    _set_spacing(p, before=0, after=after, line=1.0)
    _add_bottom_border(p, size="6")
    return p


def _add_label_with_value(doc, label: str, value: str = "", size_pt=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_spacing(p, before=0, after=3, line=1.0)

    r1 = p.add_run(label)
    _apply_font(r1, size_pt=size_pt)

    if value:
        r2 = p.add_run(value)
        _apply_font(r2, size_pt=size_pt)

    return p


def _add_multiline_field(doc, label: str, content: str = "", total_lines=3, size_pt=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_spacing(p, before=0, after=2, line=1.0)
    r = p.add_run(label)
    _apply_font(r, size_pt=size_pt)

    content = _normalize_text(content)
    lines = [x.strip() for x in content.split("\n") if x.strip()] if content else []

    for i in range(total_lines):
        pl = doc.add_paragraph()
        pl.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_spacing(pl, before=0, after=3, line=1.0)
        if i < len(lines):
            rr = pl.add_run(lines[i])
            _apply_font(rr, size_pt=size_pt)
        _add_bottom_border(pl, size="6")


def _add_signature_block(doc, text_below: str):
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing(p1, before=18, after=1, line=1.0)
    r1 = p1.add_run("________________________________________")
    _apply_font(r1, size_pt=11)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing(p2, before=0, after=8, line=1.0)
    r2 = p2.add_run(text_below)
    _apply_font(r2, size_pt=10, bold=True)


# =========================
# DOCUMENTO
# =========================

async def registro_asesoria_individual_maker(
    analysis_json: str,
    tool_context: CallbackContext,
) -> dict:
    try:
        analysis = json.loads(analysis_json) if isinstance(analysis_json, str) else analysis_json

        nombre_usuario = _safe(analysis.get("nombre_usuario"))
        dui = _safe(analysis.get("dui"))
        genero = _safe(analysis.get("genero"))
        problema_planteado = _normalize_text(analysis.get("problema_planteado"))
        asesoria_para_juicio_proceso = _normalize_text(analysis.get("asesoria_para_juicio_proceso"))
        documentos_requeridos = _normalize_text(analysis.get("documentos_requeridos"))
        requisitos_faltantes = _normalize_text(analysis.get("requisitos_faltantes"))

        doc = Document()

        section = doc.sections[0]
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

        _set_default_font(doc, font_name="Arial", size_pt=11)

        # Código superior derecho
        p_code = doc.add_paragraph()
        p_code.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_spacing(p_code, before=0, after=2, line=1.0)
        r_code = p_code.add_run("FOLA01")
        _apply_font(r_code, size_pt=9, bold=True)

        # Encabezado
        p_header = doc.add_paragraph()
        p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_spacing(p_header, before=0, after=2, line=1.0)

        for line in [
            "PROCURADURÍA GENERAL DE LA REPÚBLICA",
            "UNIDAD DE DEFENSA DE LOS DERECHOS DEL TRABAJADOR",
        ]:
            r = p_header.add_run(line + "\n")
            _apply_font(r, size_pt=11, bold=True)

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_spacing(p_title, before=8, after=12, line=1.0)
        r_title = p_title.add_run("REGISTRO DE ASESORÍA INDIVIDUAL")
        _apply_font(r_title, size_pt=14, bold=True)

        # Línea de procuraduría / hora / fecha: vacía para llenado a mano
        _add_label_with_value(
            doc,
            "Procuraduría Auxiliar de ______________________________ a las ______________________________ horas"
        )
        _add_label_with_value(
            doc,
            "____________ minutos del día __________________ de __________________ 20______."
        )

        # Nombre
        _add_label_with_value(doc, "Nombre del/la usuario/a: ", nombre_usuario)
        _add_empty_line(doc)

        # Género / edad / DUI
        _add_label_with_value(
            doc,
            "Género ",
            f"{genero} de ______ años de edad, con Documento Único de Identidad Número {dui}"
        )

        # Lugar y fecha de expedición del DUI: vacío
        _add_label_with_value(
            doc,
            "Extendido en ________________________, el día __________________ mes __________________ 20______"
        )

        # Problema planteado
        _add_multiline_field(
            doc,
            "Problema planteado: ",
            problema_planteado,
            total_lines=3
        )

        # Asesoría para juicio/proceso
        _add_multiline_field(
            doc,
            "Asesoría para Juicio / Proceso: ",
            asesoria_para_juicio_proceso,
            total_lines=2
        )

        # Documentos requeridos
        _add_multiline_field(
            doc,
            "Documentos requeridos: ",
            documentos_requeridos,
            total_lines=3
        )

        # Requisitos faltantes
        _add_multiline_field(
            doc,
            "Requisitos que faltan para conceder la asistencia legal: ",
            requisitos_faltantes,
            total_lines=4
        )

        # Firmas
        _add_signature_block(doc, "Firma o huella del usuario/a")
        _add_signature_block(doc, "Nombre y firma Defensor/a Público Laboral")

        # Footer
        footer_para = doc.sections[0].footer.paragraphs[0]
        footer_para.text = f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if footer_para.runs:
            _apply_font(
                footer_para.runs[0],
                size_pt=8,
                color_rgb=(110, 110, 110)
            )

        # Guardar a bytes
        buffer = BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.seek(0)

        mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        safe_dui = "".join(ch for ch in str(dui) if ch.isalnum()) or "sin_dui"
        case_id = uuid.uuid4().hex
        output_filename = f"{case_id}_{safe_dui}_registro_asesoria_individual.docx"
        artifact_filename = _build_artifact_filename(case_id, safe_dui, "rai")

        # GCS
        bucket_name = os.environ["GCS_BUCKET_LABORAL"]
        object_name = f"case_{safe_dui}/fola_document/{output_filename}"
        save_gcp = upload_to_gcs(
            bucket_name=bucket_name,
            object_name=object_name,
            data=docx_bytes,
        )

        artifact_part = types.Part(
            inline_data=types.Blob(data=docx_bytes, mime_type=mime_type)
        )

        version = await tool_context.save_artifact(
            filename=artifact_filename,
            artifact=artifact_part
        )

        return {
            "status": "ok",
            "message": f"El documento {output_filename} versión {version} ha sido creado.",
            "artifact_filename": artifact_filename,
            # "gcs": save_gcp
        }

    except Exception as e:
        return {
            "status": "error",
            "error_type": type(e).__name__,
            "error_message": str(e),
            "traceback": traceback.format_exc()
        }