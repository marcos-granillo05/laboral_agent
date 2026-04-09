import json
import os
import uuid
import traceback
from io import BytesIO
from datetime import datetime, timezone

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from google.cloud import storage
from google.oauth2 import service_account
from google.genai import types
from google.adk.agents.callback_context import CallbackContext


def build_storage_client():
    key_path = os.environ.get("ADK_CREDENTIALS_PATH")

    if key_path and os.path.exists(key_path):
        creds = service_account.Credentials.from_service_account_file(key_path)
        client = storage.Client(credentials=creds, project=creds.project_id)
        return client, creds

    client = storage.Client()
    return client, None


def upload_to_gcs(bucket_name: str, object_name: str, data: bytes) -> dict:
    client, _ = build_storage_client()
    bucket = client.bucket(bucket_name)
    blob = bucket.blob(object_name)

    blob.upload_from_string(
        data,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    print(f"DEBUG: upload ok bucket={bucket_name} object={object_name}")

    return {
        "bucket_name": bucket_name,
        "object_name": object_name,
    }


def _apply_font(run, font_name="Arial", size_pt=None, color_rgb=(0, 0, 0), bold=False):
    run.font.name = font_name

    rPr = run._element.get_or_add_rPr()

    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)

    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)

    run.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def _set_default_font(doc: Document, font_name="Arial", size_pt=12):
    style = doc.styles['Normal']
    style.font.name = font_name
    style.font.size = Pt(size_pt)

    rPr = style.element.get_or_add_rPr()
    rfonts = rPr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rPr.append(rfonts)

    rfonts.set(qn('w:ascii'), font_name)
    rfonts.set(qn('w:hAnsi'), font_name)
    rfonts.set(qn('w:cs'), font_name)


def _add_horizontal_rule(paragraph, size="18", color="000000", space="1"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)

    pBdr.append(bottom)
    pPr.append(pBdr)


def _fmt_fecha_from_ts(ts):
    if ts is None:
        return None

    if isinstance(ts, str):
        try:
            dt = datetime.fromisoformat(ts.replace("Z", "+00:00"))
        except ValueError:
            ts_num = float(ts)
            dt = _dt_from_epoch(ts_num)
    else:
        dt = _dt_from_epoch(float(ts))

    meses = [
        "enero", "febrero", "marzo", "abril", "mayo", "junio",
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"
    ]
    return f"{dt.day} de {meses[dt.month-1]} de {dt.year}"


def _dt_from_epoch(ts_num: float) -> datetime:
    if ts_num > 1e12:
        ts_num = ts_num / 1000.0

    return datetime.fromtimestamp(ts_num, tz=timezone.utc)


def _normalize_content(value):
    if value is None:
        return ""
    if isinstance(value, list):
        return [str(x) for x in value if str(x).strip()]
    if isinstance(value, str):
        return value.strip()

    return str(value).strip()


def _add_section(doc: Document, title: str, content, font_name="Arial"):
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = h.add_run(title)
    _apply_font(r, font_name=font_name, size_pt=12, bold=True)

    content_norm = _normalize_content(content)

    if isinstance(content_norm, list):
        for item in content_norm:
            p = doc.add_paragraph(item, style="List Bullet")
            if p.runs:
                _apply_font(p.runs[0], font_name=font_name, size_pt=11)
    else:
        chunks = [c.strip() for c in content_norm.split("\n\n") if c.strip()]
        if not chunks:
            chunks = [""]

        for ch in chunks:
            p = doc.add_paragraph(ch)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if p.runs:
                _apply_font(p.runs[0], font_name=font_name, size_pt=11)

    doc.add_paragraph()


async def document_maker(analysis_json: str, tool_context: CallbackContext) -> dict:
    try:
        analysis = json.loads(analysis_json) if isinstance(analysis_json, str) else analysis_json

        doc = Document()

        section = doc.sections[0]
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        font_name = "Arial"
        _set_default_font(doc, font_name=font_name, size_pt=12)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)

        for line in [
            "PROCURADURÍA GENERAL DE LA REPÚBLICA",
            "UNIDAD PARA LA DEFENSA DE LOS DERECHOS",
            "LABORALES DE LA PERSONA TRABAJADORA",
            "PROCURADURÍA AUXILIAR DE xxxxxxxxxx",
            "Teléfono xxxx-xxxx",
        ]:
            r = p.add_run(line + ("\n" if line != "Teléfono xxxx-xxxx" else ""))
            _apply_font(r, font_name=font_name, size_pt=11, bold=True)

        line_p = doc.add_paragraph()
        _add_horizontal_rule(line_p, size="18")
        line_p.paragraph_format.space_after = Pt(12)

        fecha_str = _fmt_fecha_from_ts(analysis.get("timestamp"))

        if fecha_str:
            pf = doc.add_paragraph(fecha_str)
            pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if pf.runs:
                _apply_font(pf.runs[0], font_name=font_name, size_pt=11)
            pf.paragraph_format.space_after = Pt(16)

        personal_data = analysis.get("personal_data_summary", "")
        summary = analysis.get("summary", "")
        justification = analysis.get("justification", "")
        claims_statement_facts = analysis.get("claims_statement_facts", "")
        user_dui = analysis.get("personal_data", {}).get("dui", "")

        _add_section(doc, "Datos Personales", personal_data, font_name=font_name)
        _add_section(doc, "I. Resumen del relato del hecho", summary, font_name=font_name)
        _add_section(doc, "II. Trámite Asociado", claims_statement_facts, font_name=font_name)
        _add_section(doc, "III. Análisis Jurídico", justification, font_name=font_name)

        atent = doc.add_paragraph("Atentamente,")
        atent.paragraph_format.space_after = Pt(28)
        if atent.runs:
            _apply_font(atent.runs[0], font_name=font_name, size_pt=11)

        firma = doc.add_paragraph()
        firma.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = firma.add_run("______________________________\n")
        _apply_font(r1, font_name=font_name, size_pt=11)
        r2 = firma.add_run("")
        _apply_font(r2, font_name=font_name, size_pt=11)

        footer_para = doc.sections[0].footer.paragraphs[0]
        footer_text = f'Generado el {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}'
        footer_para.text = footer_text
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if footer_para.runs:
            _apply_font(footer_para.runs[0], font_name=font_name, size_pt=8, color_rgb=(128, 128, 128))

        buffer = BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.seek(0)

        bucket_name = os.environ["GCS_BUCKET_LABORAL"]

        case_id = uuid.uuid4().hex
        safe_dui = "".join(ch for ch in str(user_dui) if ch.isalnum()) or "sin_dui"

        object_name = f"case_{safe_dui}/analytical_summary/{case_id}_{safe_dui}_analytical_summary_document.docx"
        output_filename = f"{case_id}_{safe_dui}_analytical_summary_document.docx"

        save_gcp = upload_to_gcs(
            bucket_name=bucket_name,
            object_name=object_name,
            data=docx_bytes,
        )

        artifact_part = types.Part(
            inline_data=types.Blob(data=docx_bytes, mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        )

        version = await tool_context.save_artifact(
            filename=output_filename,
            artifact=artifact_part
        )

        return {
            "status": "ok",
            "message": f"El documento {output_filename} version {version} ha sido creado y disponible para descargar.",
            "gcs": save_gcp
        }

    except Exception as e:
        return {
            "status": "error",
            "error_type": type(e).__name__,
            "error_message": str(e),
            "traceback": traceback.format_exc()
        }