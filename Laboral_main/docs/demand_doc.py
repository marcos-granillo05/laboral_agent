import json
import os
import uuid
import traceback
from io import BytesIO
from datetime import datetime, timezone

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from google.cloud import storage
from google.oauth2 import service_account
from google.genai import types
from google.adk.agents.callback_context import CallbackContext


def build_storage_client():
    key_path = os.environ.get("ADK_CREDENTIALS_PATH")

    if key_path and os.path.exists(key_path):
        creds = service_account.Credentials.from_service_account_file(key_path)
        client = storage.Client(credentials=creds, project=creds.project_id)
        return client, creds

    client = storage.Client()
    return client, None


def upload_to_gcs(bucket_name: str, object_name: str, data: bytes) -> dict:
    client, _ = build_storage_client()
    bucket = client.bucket(bucket_name)
    blob = bucket.blob(object_name)

    blob.upload_from_string(
        data,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    print(f"DEBUG: upload ok bucket={bucket_name} object={object_name}")

    return {
        "bucket_name": bucket_name,
        "object_name": object_name,
    }


def _build_artifact_filename(case_id: str, safe_dui: str, prefix: str) -> str:
    short_case_id = case_id[:8]
    short_dui = safe_dui[:12]
    return f"{prefix}_{short_dui}_{short_case_id}.docx"


def _apply_font(run, font_name="Arial", size_pt=12, color_rgb=(0, 0, 0), bold=False, italic=False):
    run.font.name = font_name

    rPr = run._element.get_or_add_rPr()

    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)

    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:cs"), font_name)

    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    run.font.size = Pt(size_pt)


def _set_default_font(doc: Document, font_name="Arial", size_pt=12):
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(size_pt)

    rPr = style.element.get_or_add_rPr()
    rfonts = rPr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rPr.append(rfonts)

    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:cs"), font_name)


def _configure_paragraph(
    paragraph,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line_indent_cm=1.25,
    space_before_pt=0,
    space_after_pt=0,
    line_spacing=1.0,
):
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Cm(first_line_indent_cm)
    paragraph.paragraph_format.space_before = Pt(space_before_pt)
    paragraph.paragraph_format.space_after = Pt(space_after_pt)
    paragraph.paragraph_format.line_spacing = line_spacing


def _add_run(paragraph, text, *, bold=False, color_rgb=(0, 0, 0), uppercase=False, font_name="Arial", size_pt=12):
    value = "" if text is None else str(text)
    if uppercase:
        value = value.upper()

    run = paragraph.add_run(value)
    _apply_font(run, font_name=font_name, size_pt=size_pt, color_rgb=color_rgb, bold=bold)
    return run


def _add_blank_line(doc: Document):
    p = doc.add_paragraph()
    _configure_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent_cm=0)
    return p


def _heading(doc: Document, text: str, font_name="Arial"):
    p = doc.add_paragraph()
    _configure_paragraph(
        p,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent_cm=0,
        space_before_pt=0,
        space_after_pt=0,
        line_spacing=1.0,
    )
    _add_run(p, text, bold=True, uppercase=True, font_name=font_name, size_pt=12)
    return p


def _body(doc: Document, font_name="Arial"):
    p = doc.add_paragraph()
    _configure_paragraph(
        p,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent_cm=1.25,
        space_before_pt=0,
        space_after_pt=0,
        line_spacing=1.0,
    )
    return p


def _safe_text(value, default=""):
    if value is None:
        return default
    return str(value).strip()


def _list_or_empty(value):
    if value is None:
        return []
    if isinstance(value, list):
        return value
    return [value]


async def complaint_maker(analysis_json: str, tool_context: CallbackContext) -> dict:
    try:
        analysis = json.loads(analysis_json) if isinstance(analysis_json, str) else analysis_json

        worker_information = analysis.get("worker_information", {}) or {}
        lawyer_information = analysis.get("lawyer_information", {}) or {}
        employment_relationship_data = analysis.get("employment_relationship_data", {}) or {}
        suggestions = _safe_text(analysis.get("law_suggestions"))

        doc = Document()

        section = doc.sections[0]
        section.page_width = Cm(21.59)
        section.page_height = Cm(27.94)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(3.0)

        font_name = "Arial"
        _set_default_font(doc, font_name=font_name, size_pt=12)

        # =========================
        # DATOS
        # =========================
        abogado_nombre = _safe_text(lawyer_information.get("lawyer_name"), "DIEGO FRANCISCO BRIZUELA HUEZO")
        abogado_domicilio = _safe_text(lawyer_information.get("lawyer_home"))
        abogado_dui = _safe_text(lawyer_information.get("lawyer_dui"))
        abogado_tarjeta = _safe_text(lawyer_information.get("lawyer_card"))
        abogado_edad = _safe_text(lawyer_information.get("lawyer_age"))

        trabajador_nombre = _safe_text(worker_information.get("worker_name"))
        trabajador_edad = _safe_text(worker_information.get("worker_age"))
        trabajador_estado_familiar = _safe_text(worker_information.get("worker_marital_status"))
        trabajador_profesion = _safe_text(worker_information.get("worker_profession_or_trade"))
        trabajador_nacionalidad = _safe_text(worker_information.get("worker_nationality"))
        trabajador_domicilio = _safe_text(worker_information.get("worker_address"))
        trabajador_dui = _safe_text(worker_information.get("worker_dui"))

        empresa_demandada = _safe_text(employment_relationship_data.get("company_defendant"))
        empresa_domicilio = _safe_text(employment_relationship_data.get("company_address"))
        representante_legal_nombre = _safe_text(employment_relationship_data.get("legal_representative_name"))
        representante_legal_domicilio = _safe_text(employment_relationship_data.get("legal_representative_address"))
        direccion_notificacion_empresa = _safe_text(employment_relationship_data.get("company_notification_address"))

        fecha_ingreso_texto = _safe_text(employment_relationship_data.get("employment_start_date_text"))
        cargo_nominal = _safe_text(employment_relationship_data.get("job_title"))
        lugar_trabajo = _safe_text(employment_relationship_data.get("workplace"))
        funciones_reales = _safe_text(employment_relationship_data.get("actual_functions"))
        jornada_descripcion = _safe_text(employment_relationship_data.get("workday_description"))
        horario_trabajo = _safe_text(employment_relationship_data.get("work_schedule"))
        salario_texto = _safe_text(employment_relationship_data.get("salary_text"))
        forma_pago = _safe_text(employment_relationship_data.get("payment_method"))

        fecha_despido_texto = _safe_text(employment_relationship_data.get("dismissal_date_text"))
        hora_despido_texto = _safe_text(employment_relationship_data.get("dismissal_time_text"))
        nombre_quien_despide = _safe_text(employment_relationship_data.get("person_who_dismissed_name"))
        cargo_quien_despide = _safe_text(employment_relationship_data.get("person_who_dismissed_position"))
        lugar_despido = _safe_text(employment_relationship_data.get("dismissal_place"))
        relato_hechos_adicional = _safe_text(analysis.get("relato_hechos_adicional"))

        # =========================
        # DOCUMENTO
        # =========================
        p = doc.add_paragraph()
        _configure_paragraph(
            p,
            align=WD_ALIGN_PARAGRAPH.LEFT,
            first_line_indent_cm=0,
            space_before_pt=0,
            space_after_pt=0,
            line_spacing=1.0,
        )
        _add_run(p, "SEÑOR JUEZ DE LO LABORAL", bold=True, uppercase=True, font_name=font_name, size_pt=12)

        _add_blank_line(doc)

        p = _body(doc, font_name=font_name)
        _add_run(p, abogado_nombre, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, ", de ", font_name=font_name)
        _add_run(p, abogado_edad, uppercase=True, font_name=font_name)
        _add_run(p, " años de edad, Abogado y del domicilio de ", font_name=font_name)
        _add_run(p, abogado_domicilio, uppercase=True, font_name=font_name)
        _add_run(p, " con Documento Único de Identidad número ", font_name=font_name)
        _add_run(p, abogado_dui, uppercase=True, font_name=font_name)
        _add_run(p, ", Tarjeta de Identificación de la Abogacía ", font_name=font_name)
        _add_run(p, abogado_tarjeta, uppercase=True, font_name=font_name)
        _add_run(p, ", a usted atentamente ", font_name=font_name)
        _add_run(p, "EXPONGO", bold=True, uppercase=True, font_name=font_name)
        _add_run(p, ":", font_name=font_name)

        _add_blank_line(doc)

        _heading(doc, "PARTE EXPOSITIVA", font_name=font_name)

        p = _body(doc, font_name=font_name)
        _add_run(p, "En mi calidad de Defensor Publico Laboral, vengo a promover Proceso Abreviado Laboral, en nombre y representación del trabajador ", font_name=font_name)
        _add_run(p, trabajador_nombre, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, ", de ", font_name=font_name)
        _add_run(p, trabajador_edad, uppercase=True, font_name=font_name)
        _add_run(p, " años de edad, ", font_name=font_name)
        _add_run(p, trabajador_estado_familiar, uppercase=True, font_name=font_name)
        _add_run(p, ", ", font_name=font_name)
        _add_run(p, trabajador_profesion, uppercase=True, font_name=font_name)
        _add_run(p, ", de nacionalidad ", font_name=font_name)
        _add_run(p, trabajador_nacionalidad, uppercase=True, font_name=font_name)
        _add_run(p, " y del domicilio de ", font_name=font_name)
        _add_run(p, trabajador_domicilio, uppercase=True, font_name=font_name)
        _add_run(p, ", con Documento Único de Identidad número ", font_name=font_name)
        _add_run(p, trabajador_dui, uppercase=True, font_name=font_name)
        _add_run(p, " contra ", font_name=font_name)
        _add_run(p, empresa_demandada, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, ", y del domicilio de ", font_name=font_name)
        _add_run(p, empresa_domicilio, uppercase=True, font_name=font_name)
        _add_run(p, ", representada legalmente por ", font_name=font_name)
        _add_run(p, representante_legal_nombre, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, ", mayor de edad y del domicilio de ", font_name=font_name)
        _add_run(p, representante_legal_domicilio, uppercase=True, color_rgb=(255, 0, 0), font_name=font_name)
        _add_run(p, ", pudiendo ser citada, notificada y emplazada dicha persona jurídica por medio de su representante legal en ", font_name=font_name)
        _add_run(p, direccion_notificacion_empresa, uppercase=True, color_rgb=(255, 0, 0), font_name=font_name)
        _add_run(p, ", para mejor ubicación agrego croquis, lugar donde habitualmente atiende sus negocios, para reclamarle prestaciones laborales.", font_name=font_name)

        _heading(doc, "RELACION DE TRABAJO", font_name=font_name)

        p = _body(doc, font_name=font_name)
        _add_run(p, "Mi representado ingresó a laborar para y a las órdenes de ", font_name=font_name)
        _add_run(p, empresa_demandada, uppercase=True, font_name=font_name)
        _add_run(p, ", el ", font_name=font_name)
        _add_run(p, fecha_ingreso_texto, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, ", con el cargo de ", font_name=font_name)
        _add_run(p, cargo_nominal, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, "; desarrollando sus labores en ", font_name=font_name)
        _add_run(p, lugar_trabajo, uppercase=True, font_name=font_name)
        _add_run(p, " y las cuales consistían en ", font_name=font_name)
        _add_run(p, funciones_reales, uppercase=True, font_name=font_name)
        _add_run(p, "; estando sujeto a una jornada ordinaria de trabajo de ", font_name=font_name)
        _add_run(p, jornada_descripcion, uppercase=True, font_name=font_name)
        _add_run(p, ", y un horario de trabajo ", font_name=font_name)
        _add_run(p, horario_trabajo, uppercase=True, font_name=font_name)
        _add_run(p, "; devengando por sus servicios un salario de ", font_name=font_name)
        _add_run(p, salario_texto, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, ", los cuales eran cancelados ", font_name=font_name)
        _add_run(p, forma_pago, uppercase=True, font_name=font_name)
        _add_run(p, ".", font_name=font_name)

        _heading(doc, "RELACION DE HECHOS", font_name=font_name)

        p = _body(doc, font_name=font_name)
        _add_run(p, "En las condiciones de trabajo antes mencionadas laboró mi patrocinado para y a las órdenes de ", font_name=font_name)
        _add_run(p, empresa_demandada, uppercase=True, font_name=font_name)
        _add_run(p, ", desde la fecha de su ingreso hasta el ", font_name=font_name)
        _add_run(p, fecha_despido_texto, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, ", fecha en la cual como a eso de las ", font_name=font_name)
        _add_run(p, hora_despido_texto, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, " ", font_name=font_name)
        _add_run(p, nombre_quien_despide, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, ", ", font_name=font_name)
        _add_run(p, cargo_quien_despide, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, ", quien tiene facultades para contratar, despedir, dirigir y administrar trabajadores, le manifestó que a partir de ese momento estaba despedido de su trabajo, hecho que ocurrió en ", font_name=font_name)
        _add_run(p, lugar_despido, uppercase=True, font_name=font_name)
        _add_run(p, ".", font_name=font_name)

        if relato_hechos_adicional:
            _add_run(p, " ", font_name=font_name)
            _add_run(p, relato_hechos_adicional, font_name=font_name)

        _add_blank_line(doc)

        _heading(doc, "PARTE PETITORIA", font_name=font_name)

        p = _body(doc, font_name=font_name)
        _add_run(p, "Por lo antes expuesto, a usted respetuosamente PIDO:", font_name=font_name)

        _add_blank_line(doc)

        p = _body(doc, font_name=font_name)
        _add_run(p, "Me admita la presente demanda, me tenga por parte en la calidad en que comparezco, cite a conciliación a ", font_name=font_name)
        _add_run(p, empresa_demandada, bold=True, uppercase=True, font_name=font_name)
        _add_run(p, " por medio de su representante legal antes mencionado, y si no llegásemos a ningún acuerdo en dicha audiencia, previo los trámites legales y las pruebas que oportunamente aportaré, sea condenada en la sentencia definitiva a pagarle a mi representado:", font_name=font_name)

        if suggestions:
            raw_items = [x.strip() for x in suggestions.split("),") if x.strip()]
            items = []

            for raw in raw_items:
                text = raw
                if not text.endswith(")"):
                    text += ")"
                items.append(text)

            for item in items:
                bp = doc.add_paragraph(style="List Bullet")
                bp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(0)
                bp.paragraph_format.line_spacing = 1.0

                r = bp.add_run(item)
                _apply_font(r, font_name=font_name, size_pt=12)

        _add_blank_line(doc)

        # Estáticos por ahora
        p = _body(doc, font_name=font_name)
        _add_run(
            p,
            "Legitimo mi personería con copia certificada por notario de la Credencial Única y sus respectivas copias de ley.",
            font_name=font_name,
        )

        _add_blank_line(doc)

        p = _body(doc, font_name=font_name)
        _add_run(p, "Señalo para oír notificaciones, ", font_name=font_name)
        _add_run(p, "uddt.sansalvador@pgr.gob.sv", font_name=font_name)

        _add_blank_line(doc)

        p = _body(doc, font_name=font_name)
        _add_run(
            p,
            "Distrito de San Salvador, Municipio de San Salvador Centro, Departamento de San Salvador, veintiuno de enero de dos mil veintiséis.",
            font_name=font_name,
        )

        # Footer
        footer_para = doc.sections[0].footer.paragraphs[0]
        footer_text = f'Generado el {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}'
        footer_para.text = footer_text
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if footer_para.runs:
            _apply_font(footer_para.runs[0], font_name=font_name, size_pt=8, color_rgb=(128, 128, 128))

        # Guardar en memoria
        buffer = BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.seek(0)

        bucket_name = os.environ["GCS_BUCKET_LABORAL"]

        case_id = uuid.uuid4().hex
        safe_dui = "".join(ch for ch in str(trabajador_dui) if ch.isalnum()) or "sin_dui"

        object_name = f"case_{safe_dui}/demandas/{case_id}_{safe_dui}_demanda_despido_directo.docx"
        output_filename = f"{case_id}_{safe_dui}_demanda_despido_directo.docx"
        artifact_filename = _build_artifact_filename(case_id, safe_dui, "ddd")

        save_gcp = upload_to_gcs(
            bucket_name=bucket_name,
            object_name=object_name,
            data=docx_bytes,
        )

        artifact_part = types.Part(
            inline_data=types.Blob(
                data=docx_bytes,
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        )

        version = await tool_context.save_artifact(
            filename=artifact_filename,
            artifact=artifact_part
        )

        return {
            "status": "ok",
            "message": f"El documento {output_filename} version {version} ha sido creado y disponible para descargar.",
            "gcs": save_gcp,
            "artifact_filename": artifact_filename,
        }

    except Exception as e:
        return {
            "status": "error",
            "error_type": type(e).__name__,
            "error_message": str(e),
            "traceback": traceback.format_exc(),
        }
